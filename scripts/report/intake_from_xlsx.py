#!/usr/bin/env python
"""
Turn a filled client intake form (XLSX) into the YAML the bank generator reads.

Why this exists: the client fills a spreadsheet, and without this step somebody
retypes 42 fields into YAML by hand - which is exactly where the category terms
get quietly changed and the bank stops being the client's own words. This reads
the sheet the client returned, writes clients/<client>.yaml, and refuses to
write when a field that measurably costs questions is empty.

Usage (from the repository root):
    python scripts/report/intake_from_xlsx.py ..\\clients\\acme-form.xlsx
    python scripts/report/intake_from_xlsx.py ..\\clients\\acme-form.xlsx --build

    --out=DIR    where the YAML goes. Defaults to <workspace>/clients, which is
                 OUTSIDE this repository: the filled form names the client,
                 their competitors and their forbidden topics.
    --date=DATE  generation date used in the filename; defaults to today.
    --build      also run scripts/report/build-question-bank.mts on the result.
"""

from __future__ import annotations

import argparse
import datetime as dt
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve()
REPO = HERE.parent.parent.parent           # geo-scanner/
WORKSPACE = REPO.parent                    # the directory clients/ lives in

SHEET = "填写表"
HEADER_ROW = 3

# Looked up by header text, not by position: a client may insert a column, and a
# 飞书多维表格 export carries the same names but not necessarily the same order.
COL_FILL = "请在这里填写"
COL_KEY = "机器字段"
COL_TYPE = "类型"
COL_OPTIONS = "选项"

# Measured 2026-10-08: emptying any of these three drops a whole group below its
# target (facts 3/7, category 2/8, comparison 4/6). Every other field can be
# blank without changing the question count - it narrows coverage instead.
HARD_FIELDS = {
    "brand.name": "事实核验那一组会掉到 3/7",
    "industry.category_terms": "品类发现那一组会掉到 2/8",
    "competitors.names": "对比选型那一组会掉到 4/6",
}


def cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def split_list(value: str) -> list[str]:
    """One item per line. Tolerates bullets and blank lines; keeps the client's order."""
    out = []
    for raw in value.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        item = raw.strip().lstrip("-•·*").strip()
        if item and not item.startswith("（"):
            out.append(item)
    return out


def quote(text: str) -> str:
    """A double-quoted YAML scalar. Sufficient for the text these forms carry."""
    body = text.replace("\\", "\\\\").replace('"', '\\"')
    body = body.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\\n")
    return f'"{body}"'


def emit(node: dict, indent: int = 0) -> list[str]:
    pad = "  " * indent
    lines: list[str] = []
    for key, value in node.items():
        if isinstance(value, dict):
            lines.append(f"{pad}{key}:")
            lines.extend(emit(value, indent + 1))
        elif isinstance(value, list):
            if not value:
                lines.append(f"{pad}{key}: []")
            else:
                lines.append(f"{pad}{key}:")
                lines.extend(f"{pad}  - {quote(item)}" for item in value)
        elif value == "":
            lines.append(f"{pad}{key}:")
        else:
            lines.append(f"{pad}{key}: {quote(value)}")
    return lines


def set_path(doc: dict, path: str, value) -> None:
    parts = path.split(".")
    node = doc
    for part in parts[:-1]:
        node = node.setdefault(part, {})
    node[parts[-1]] = value


def read_form(path: Path) -> tuple[dict, dict, list[str]]:
    wb = load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"工作表 '{SHEET}' 不在 {path.name} 里;可用:{', '.join(wb.sheetnames)}")
    ws = wb[SHEET]

    header = {cell_text(c.value): c.column for c in ws[HEADER_ROW] if cell_text(c.value)}
    wanted = {name: header.get(name) for name in (COL_FILL, COL_KEY, COL_TYPE, COL_OPTIONS)}
    if not wanted[COL_KEY] or not wanted[COL_FILL]:
        raise SystemExit(
            f"第 {HEADER_ROW} 行找不到列「{COL_KEY}」/「{COL_FILL}」。这份表必须来自 "
            "intake/_form-template.xlsx(飞书版请把列名建成一样的)。"
        )

    def at(row, name):
        col = wanted[name]
        return cell_text(row[col - 1].value) if col else ""

    doc: dict = {}
    meta: dict = {}
    problems: list[str] = []

    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        key = at(row, COL_KEY)
        if not key:
            continue
        raw = at(row, COL_FILL)
        kind = at(row, COL_TYPE)
        options = at(row, COL_OPTIONS)

        value: object
        if "(一行一条)" in kind:
            value = split_list(raw)
        else:
            value = raw

        if options and raw:
            allowed = [o.strip() for o in options.split("/")]
            if raw not in allowed:
                problems.append(f"{key}：值 {raw!r} 不在允许的取值里（{' / '.join(allowed)}）—— 这一栏是封闭选项,生成器会直接报错")

        meta[key] = {"raw": raw, "value": value, "kind": kind, "options": options}
        set_path(doc, key, value)

    return doc, meta, problems


def ascii_slug(text: str) -> str:
    """At least one ASCII character is required: the bank generator derives its output
    filename from meta.client OR the intake filename, and a pure-Chinese client name
    fails both (build-question-bank.mts: asciiSlug(client) || asciiSlug(filename)).
    So the YAML file gets a Latin slug while meta.client keeps the real name."""
    out = []
    for ch in text.lower():
        out.append(ch if ch.isalnum() and ch.isascii() else "-")
    slug = "".join(out)
    while "--" in slug:
        slug = slug.replace("--", "-")
    return slug.strip("-")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("form", type=Path, help="填好的 intake 表(.xlsx)")
    ap.add_argument("--out", type=Path, default=WORKSPACE / "clients")
    ap.add_argument("--date", default=dt.date.today().isoformat())
    ap.add_argument("--slug", default="", help="输出文件用的拉丁短名,如 champion-coatings。客户名全中文时必填")
    ap.add_argument("--approved-by", default="", help="题库确认人的名字。报告第一节会印这个名字")
    ap.add_argument("--build", action="store_true", help="生成后立刻跑题库生成器")
    args = ap.parse_args()

    if not args.form.exists():
        raise SystemExit(f"找不到文件:{args.form}")

    doc, meta, problems = read_form(args.form)
    doc.setdefault("meta", {})

    client = cell_text(meta.get("meta.client", {}).get("raw", ""))
    if not client:
        problems.append("meta.client 为空:公司名必须填,它决定文件命名")

    hard_empty = [k for k in HARD_FIELDS if not (meta.get(k, {}).get("raw") or "").strip()]
    for key in hard_empty:
        problems.append(f"{key} 为空 —— {HARD_FIELDS[key]}(实测,不是估计)")

    missing_required = [
        k for k, m in meta.items()
        if k not in HARD_FIELDS and m["kind"] and m["raw"] == "" and k != "meta.client"
    ]

    if meta.get("compliance.sensitive", {}).get("raw") == "是" and not meta.get("compliance.restrictions", {}).get("raw"):
        problems.append("compliance.sensitive = 是 但 restrictions 为空 —— 这一栏空着不建议开跑")

    doc["meta"]["filled_by"] = "intake form (xlsx)"
    doc["meta"]["filled_on"] = args.date
    # 签字是客户在确认题库时给的,不是填表时填的 —— 所以从命令行传,而不是从表单读。
    # 它会进题库的 approval 块,并由采集脚本写进 run header,报告第一节照此打印。
    doc["meta"]["bank_approved_by"] = args.approved_by

    print(f"表单    {args.form}")
    print(f"客户    {client or '（未填）'}")
    print(f"字段    {len(meta)} 行,其中硬字段 {len(HARD_FIELDS) - len(hard_empty)}/{len(HARD_FIELDS)} 已填")

    if problems:
        print("\n不能生成,先修这几处：")
        for p in problems:
            print(f"  ✗ {p}")
        return 1

    if missing_required:
        print(f"\n提示({len(missing_required)} 栏为空,题数仍是 29,但覆盖面会变窄)：")
        for k in missing_required:
            print(f"  · {k}")

    safe = args.slug.strip() or ascii_slug(client)
    if not safe:
        print(
            "\n不能生成:客户名里没有 ASCII 字符,而题库生成器需要一个能当文件名的 slug。\n"
            "  加一个拉丁短名再跑,例如:--slug=champion-coatings\n"
            "  meta.client 里仍保留完整中文名,报告与题库里显示的是它。"
        )
        return 1
    out_dir = args.out.resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{safe}-{args.date}.yaml"
    out_path.write_text("\n".join(emit(doc)) + "\n", encoding="utf-8")
    print(f"\n已写出  {out_path}")
    print(f"        slug={safe}(只用于文件名;meta.client 仍是「{client}」)")
    if REPO in out_path.parents:
        print("⚠️ 这份文件在仓库里,而它含客户机密 —— 不要提交。")

    if args.build:
        script = REPO / "scripts" / "report" / "build-question-bank.mts"
        print(f"\n跑题库生成器 …")
        code = subprocess.call(["node", str(script), f"--intake={out_path}"], cwd=REPO)
        if code != 0:
            print(f"生成器退出码 {code}")
            return code
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
